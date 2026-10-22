"""Knowledge base: turn the workbook + field_pipeline results into one fact table and a comparison page.

  python -m knowledge_base "../EDIT Competitive Analysis Research 2026-2027.xlsx"

Each fact is one (program, field) value with its cycle, status, source URL and quote. A value comes from, in order:
  1 this year's field_pipeline decisions.csv (found / found-prior / computed / carried-forward / ipeds)
  2 a value already typed into this year's workbook
  3 last year's workbook (cycle 2025-2026), matched by school name; schools only in last year's peer set are kept
Writes <workbook>_kb/kb.sqlite (for queries) and <workbook>_kb/competitive_analysis.html (open in any browser;
the data is embedded, so the file can be shared on its own). No web or LLM calls: it runs in seconds.
"""
import argparse
import csv
import json
import re
import sqlite3
from datetime import date
from pathlib import Path

from field_pipeline.fields import FILLED, LINK, SKIP, field_for, link_target
from field_pipeline.validate import NUMERIC, to_number
from field_pipeline.workbook import headers, read_rows, read_values

from openpyxl import load_workbook

GROUP_LABELS = {
    "basics": "Basics", "scores": "Test scores & GPA", "class": "Class profile", "structure": "Program structure",
    "campus": "On-campus", "deadlines": "Deadlines", "requirements": "Admissions", "tuition": "Tuition",
    "aid": "Scholarships & deposit", "experiential": "Experiential learning", "community": "Community",
    "career": "Career services", "outcomes": "Outcomes", "rank_usnews": "Rankings", "rank_bestmasters": "Rankings",
    "rank_qs": "Rankings", "other": "Other",
}
NEW = re.compile(r"\s*(\(new 2026\)|-\s*2026)\s*$", re.I)
JUNK = re.compile(r"\d+|.*not ranked.*", re.I)  # note rows in the sheets, not schools
ID_COLUMNS = re.compile(r"^(university( name)?|school\s*/\s*college( name)?|program name|website link)$", re.I)


def text(v) -> str:
    """Cell value as text; "" for blanks, placeholder dashes and Excel formulas (last year's ROI columns)."""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    v = " ".join(str(v).split())
    return "" if re.fullmatch(r"[-–—?.\s]*", v) or v.startswith("=") else v


def current_values(path: str) -> dict[tuple[str, int], dict]:
    """{(sheet, row): {column: value}} for values already typed into this year's workbook."""
    wb = load_workbook(path, read_only=True)
    out = {}
    for ws in wb.worksheets:
        head = headers(ws)
        if "Website Link" not in head:
            continue
        for r, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            out[(ws.title, r)] = {h: v for h, v in zip(head, values) if h and v not in (None, "")}
    wb.close()
    return out


def decisions(path: Path) -> dict[tuple[str, int, str], dict]:
    if not path.exists():
        return {}
    with open(path, newline="", encoding="utf-8") as f:
        return {(d["sheet"], int(d["row"]), d["field"]): d for d in csv.DictReader(f) if d["status"] in FILLED}


def words(s) -> list[str]:
    return [w for w in re.sub(r"[^a-z0-9]+", " ", str(s or "").casefold().replace("&", " and ")).split()
            if w not in ("the", "at", "of", "and")]


def school_of(rec: dict) -> str:
    return next((text(v) for c, v in rec.items() if re.match(r"school\s*/\s*college", str(c), re.I) and v), "")


def col_key(c: str) -> str:
    """Columns match across years by field (\"Female %\" = \"Female (%)\"), else by their words."""
    f = field_for(c)
    return f.pattern if f else " ".join(sorted(words(c)))


def match_prior(univ: str, program: str, recs: list[tuple[str, dict]], used: set[int]) -> int | None:
    """Index of last year's row for this school: same name, or one name starts the other ("Purdue University" ~
    "PURDUE UNIVERSITY - WEST LAFAYETTE", not "University of Washington" ~ "Washington State University");
    several -> the closest program name, ties -> none."""
    u = words(univ)
    prefix = lambda a, b: a == b[:len(a)] or b == a[:len(b)]
    free = [i for i in range(len(recs)) if i not in used]
    cands = [i for i in free if words(recs[i][0]) == u] or [i for i in free if prefix(words(recs[i][0]), u)]
    if len(cands) > 1:
        score = lambda i: len(set(words(recs[i][1].get("Program Name"))) & set(words(program)))
        best = max(map(score, cands))
        cands = [i for i in cands if score(i) == best]
    return cands[0] if len(cands) == 1 else None


def build(workbook: str, prior_path: str | None, decisions_path: Path) -> dict:
    cols, rows = read_rows(workbook)
    now, found = current_values(workbook), decisions(decisions_path)
    prior: dict[str, list[tuple[str, dict]]] = {}
    for (sheet, univ), recs in (read_values(prior_path) if prior_path else {}).items():
        if not JUNK.fullmatch(univ):
            prior.setdefault(sheet, []).extend((univ, r) for r in recs)
    sheets, programs, facts = [], [], []
    for sheet, head in cols.items():
        fields, links = [], {}
        for c in head:
            if not c or c.startswith("Unnamed") or ID_COLUMNS.match(c):
                continue
            if pair := link_target(c, head):
                links[c] = pair
                continue
            f = field_for(c)
            label = GROUP_LABELS.get(f.group if f else "other", "Other")
            fields.append({"name": c, "group": label, "kind": f.kind if f else ("number" if SKIP.search(c) else "text")})
        order = list(dict.fromkeys(GROUP_LABELS.values()))
        fields.sort(key=lambda f: order.index(f["group"]))  # workbook order within each group
        sheets.append({"name": sheet, "fields": fields})
        names = {f["name"]: f for f in fields}
        by_key = {col_key(n): n for n in names}
        old_recs, used = prior.get(sheet.casefold(), []), set()

        def last_year(rec: dict) -> dict[str, dict]:
            """{current column: fact} from a row of last year's workbook. "(New 2026)" columns were typed for 2026-27."""
            out = {}
            for c, v in sorted(rec.items(), key=lambda kv: bool(NEW.search(str(kv[0])))):
                if str(c).startswith("_") or ID_COLUMNS.match(str(c)) or LINK.search(str(c)):
                    continue
                name = by_key.get(col_key(NEW.sub("", str(c))))
                if name and text(v):
                    new = bool(NEW.search(str(c)))
                    out[name] = {"v": text(v), "c": "2026-2027" if new else "2025-2026",
                                 "s": "entered" if new else "prior-year", "u": "", "q": ""}
            for c, v in rec.items():  # last year's link columns -> source of the paired field
                if (pair := link_target(str(c), list(rec))) and (name := by_key.get(col_key(pair))) in out and text(v):
                    out[name]["u"] = text(v)
            return out

        def add(pid: int, values: dict[str, dict]) -> None:
            for name, fact in values.items():
                if not fact["v"]:
                    continue
                kind = names[name]["kind"]
                n = to_number(fact["v"], kind) if kind in NUMERIC else None
                if kind == "percent" and n is not None and n <= 1 and re.fullmatch(r"[\d.]+", fact["v"]):
                    n = round(n * 100, 1)  # Excel stores 28% as 0.28
                    fact["v"] = f"{n:g}%"
                facts.append({"p": pid, "f": name, "n": n, **fact})

        for row in (r for r in rows if r.sheet == sheet and not JUNK.fullmatch(r.university)):
            typed = now.get((sheet, row.row), {})
            i = match_prior(row.university, row.program, old_recs, used)
            old_rec = old_recs[i][1] if i is not None else {}
            if i is not None:
                used.add(i)
            old = last_year(old_rec)
            pid = len(programs)
            programs.append({"id": pid, "sheet": sheet, "university": row.university, "program": row.program,
                             "school": school_of(typed) or school_of(old_rec), "url": row.start_url, "purdue": words(row.university) in (["purdue", "university"], ["purdue", "university", "west", "lafayette"]),
                             "lastYearOnly": False})
            values = {}
            for name in names:
                link = next((text(typed[l]) for l, p in links.items() if p == name and typed.get(l)), "")
                if d := found.get((sheet, row.row, name)):
                    values[name] = {"v": text(d["value"]), "c": d["cycle"] or "2026-2027", "s": d["status"],
                                    "u": d["source_url"], "q": d["quote"]}
                    if d["status"].startswith("ipeds"):
                        values[name]["d"] = d["detail"]
                elif text(typed.get(name) or ""):
                    values[name] = {"v": text(typed[name]), "c": "2026-2027", "s": "entered", "u": link, "q": ""}
                elif name in old:
                    values[name] = old[name]
            add(pid, values)
        for i, (univ, rec) in enumerate(old_recs):  # schools only in last year's peer set
            if i in used or not (values := last_year(rec)):
                continue
            pid = len(programs)
            programs.append({"id": pid, "sheet": sheet, "university": univ.upper(), "program": text(rec.get("Program Name") or ""),
                             "school": school_of(rec), "url": text(rec.get("Website Link") or ""),
                             "purdue": False, "lastYearOnly": True})
            add(pid, values)
    return {"built": date.today().isoformat(), "sheets": sheets,
            "programs": programs, "facts": facts}


def write_sqlite(kb: dict, path: Path) -> None:
    path.unlink(missing_ok=True)
    db = sqlite3.connect(path)
    db.executescript("""
        CREATE TABLE programs (id INTEGER PRIMARY KEY, sheet TEXT, university TEXT, school TEXT, program TEXT,
                               url TEXT, purdue INTEGER, last_year_only INTEGER);
        CREATE TABLE fields (sheet TEXT, name TEXT, field_group TEXT, kind TEXT);
        CREATE TABLE facts (program_id INTEGER REFERENCES programs(id), field TEXT, value TEXT, number REAL,
                            cycle TEXT, status TEXT, source_url TEXT, quote TEXT);
        CREATE INDEX facts_field ON facts (field, number);
    """)
    db.executemany("INSERT INTO programs VALUES (:id, :sheet, :university, :school, :program, :url, :purdue, :lastYearOnly)",
                   kb["programs"])
    db.executemany("INSERT INTO fields VALUES (?, ?, ?, ?)",
                   [(s["name"], f["name"], f["group"], f["kind"]) for s in kb["sheets"] for f in s["fields"]])
    db.executemany("INSERT INTO facts VALUES (:p, :f, :v, :n, :c, :s, :u, :q)", kb["facts"])
    db.commit()
    db.close()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--prior", type=Path, help="last year's workbook (default: the one OLD *.xlsx next to it)")
    ap.add_argument("--decisions", type=Path, help="field_pipeline decisions.csv (default: <workbook>_field_pipeline/)")
    ap.add_argument("--out", type=Path, help="output folder (default: <workbook>_kb/)")
    args = ap.parse_args()
    wb = args.workbook
    olds = sorted(wb.parent.glob("OLD *.xlsx"))
    prior = args.prior or (olds[0] if len(olds) == 1 else None)
    out = args.out or wb.with_name(f"{wb.stem}_kb")
    out.mkdir(parents=True, exist_ok=True)
    kb = build(str(wb), str(prior) if prior else None,
               args.decisions or wb.with_name(f"{wb.stem}_field_pipeline") / "decisions.csv")
    write_sqlite(kb, out / "kb.sqlite")
    page = (Path(__file__).with_name("template.html").read_text(encoding="utf-8")
            .replace("/*KB_DATA*/null", json.dumps(kb, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")))
    (out / "competitive_analysis.html").write_text(page, encoding="utf-8")
    by = {}
    for f in kb["facts"]:
        by[f["s"]] = by.get(f["s"], 0) + 1
    print(f"{len(kb['programs'])} programs, {len(kb['facts'])} facts {by} -> {out}")


if __name__ == "__main__":
    main()
