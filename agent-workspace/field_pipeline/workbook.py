"""Read program rows from the workbook and write results to a copy (formatting kept, only target cells touched)."""
import re
from dataclasses import dataclass

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

AUTHOR = "field-pipeline"


@dataclass
class Row:
    sheet: str
    row: int              # 1-based Excel row
    university: str
    program: str
    start_url: str
    blanks: list[str]     # columns that are empty in the workbook

    @property
    def key(self) -> str:
        return f"{self.sheet}::{self.row}::{self.university}::{self.program}"


def _blank(v) -> bool:
    return v is None or str(v).strip() in ("", "nan", "None")


def headers(ws) -> list[str]:
    return [" ".join(str(c.value).split()) if c.value is not None else None for c in ws[1]]


def read_rows(path: str, sheets: list[str] | None = None) -> tuple[dict[str, list[str]], list[Row]]:
    wb = load_workbook(path, read_only=True)
    cols, rows = {}, []
    for ws in wb.worksheets:
        if sheets and ws.title not in sheets:
            continue
        head = headers(ws)
        if "Website Link" not in head:
            continue
        cols[ws.title] = head
        web = head.index("Website Link")
        prog = next((i for i, h in enumerate(head) if h and re.fullmatch(r"program name", h, re.I)), 2)
        for r, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            values = list(values) + [None] * (len(head) - len(values))
            if _blank(values[0]):
                continue
            rows.append(Row(ws.title, r, str(values[0]).strip(),
                            "" if _blank(values[prog]) else str(values[prog]).strip(),
                            "" if _blank(values[web]) else str(values[web]).strip(),
                            [h for h, v in zip(head, values) if h and _blank(v)]))
    wb.close()
    return cols, rows


def read_values(path: str) -> dict[tuple[str, str], list[dict]]:
    """Every program row of a workbook as {column: value}, keyed by (sheet, university) casefolded."""
    wb = load_workbook(path, read_only=True)
    out: dict[tuple[str, str], list[dict]] = {}
    for ws in wb.worksheets:
        head = headers(ws)
        if "Website Link" not in head:
            continue
        for r, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if values and not _blank(values[0]):
                rec = {h: v for h, v in zip(head, values) if h and not _blank(v)}
                rec["_row"] = r
                out.setdefault((ws.title.casefold(), str(values[0]).strip().casefold()), []).append(rec)
    wb.close()
    return out


def write(src: str, dest: str, cells: list[tuple]) -> int:
    """cells: (sheet, row, column, value or None, comment, hyperlink[, fill hex]). None value = comment only."""
    wb = load_workbook(src)
    n = 0
    for sheet, row, col, value, note, link, *fill in cells:
        ws = wb[sheet]
        head = headers(ws)
        cell = ws.cell(row=row, column=head.index(col) + 1)
        if value is not None:
            cell.value = value
            n += 1
        cell.comment = Comment(note[:1500], AUTHOR, width=420, height=180)
        if link and value is not None:
            cell.hyperlink = link
        if fill and fill[0] and value is not None:
            cell.fill = PatternFill("solid", start_color=fill[0], end_color=fill[0])
    wb.save(dest)
    return n
